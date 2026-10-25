"""课程包自检：文件齐全、课件完整、链接可达、同步状态正常。

用法：python tools/verify_course.py
退出码：0 全部通过；1 有错误（错误会逐条打印）
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

REQUIRED_DOCS = [
    "README.md", "00_课程总览.md", "01_课程标准.md", "02_课程授课计划.md",
    "04_考核评价与评分量表.md", "05_实训环境与教学资源.md", "06_双元制本土化实施指南.md",
    "07_教师手册.md", "08_学生手册.md", "10_题库与参考答案.md", "11_项目案例库与代码导览.md",
    "PROJECT_UPDATES.md", "AGENTS.md",
]
REQUIRED_PLANS = [
    "03_课程教学方案/00_教学方案总则.md",
    "03_课程教学方案/01_模块一_认知与规格.md",
    "03_课程教学方案/02_模块二_智能体内核.md",
    "03_课程教学方案/03_模块三_决策智能.md",
    "03_课程教学方案/04_模块四_工程化交付.md",
]

errors: list[str] = []
warnings: list[str] = []


def check_docs() -> None:
    for rel in REQUIRED_DOCS + REQUIRED_PLANS:
        p = ROOT / rel
        if not p.exists():
            errors.append(f"缺少文件：{rel}")
        elif p.stat().st_size < 500:
            errors.append(f"文件过小（可能为空）：{rel}")


def check_taskbooks() -> None:
    files = sorted((ROOT / "09_项目任务书").glob("T*.md"))
    if len(files) != 16:
        errors.append(f"任务书应为 16 份，实际 {len(files)} 份")
    for i in range(1, 17):
        if not any(f.name.startswith(f"T{i:02d}") for f in files):
            errors.append(f"缺少任务书 T{i:02d}")


def check_slides() -> None:
    try:
        from pptx import Presentation
    except ImportError:
        warnings.append("未安装 python-pptx，跳过课件检查")
        return
    decks = sorted((ROOT / "slides" / "pptx").glob("*.pptx"))
    if len(decks) != 16:
        errors.append(f"课件应为 16 套，实际 {len(decks)} 套")
    for i in range(1, 17):
        matches = [d for d in decks if d.name.startswith(f"第{i:02d}讲")]
        if not matches:
            errors.append(f"缺少第 {i:02d} 讲课件")
            continue
        prs = Presentation(str(matches[0]))
        n = len(prs.slides)
        if n < 14:
            errors.append(f"{matches[0].name} 页数过少：{n}")
        noted = 0
        for s in prs.slides:
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
                noted += 1
        if noted / max(n, 1) < 0.7:
            warnings.append(f"{matches[0].name} 讲稿备注覆盖不足：{noted}/{n}")
    print(f"  课件：{len(decks)} 套")


def check_links() -> None:
    pattern = re.compile(r"\[[^\]]*\]\((?!https?://|mailto:|#)([^)\s]+)\)")
    bad = 0
    for md in ROOT.rglob("*.md"):
        if ".qa" in md.parts or ".git" in md.parts:
            continue
        text = md.read_text(encoding="utf-8", errors="replace")
        for m in pattern.finditer(text):
            target = m.group(1).split("#")[0]
            if not target or target.startswith("http"):
                continue
            if not (md.parent / target).exists():
                bad += 1
                if bad <= 10:
                    errors.append(f"链接失效：{md.relative_to(ROOT)} → {target}")
    if bad:
        errors.append(f"共 {bad} 个失效链接")


def check_sync() -> None:
    state = ROOT / "sync" / "upstream-state.json"
    if not state.exists():
        warnings.append("尚未建立上游同步基线（运行 tools/sync_course.py --init）")
        return
    data = json.loads(state.read_text(encoding="utf-8"))
    if not data.get("last_commit"):
        warnings.append("同步基线缺少 commit")
    print(f"  同步基线：{data.get('last_commit', '')[:7]} @ {data.get('last_sync_at', '')}")


def main() -> None:
    print("课程包自检")
    check_docs()
    check_taskbooks()
    check_slides()
    check_links()
    check_sync()

    if warnings:
        print("\n警告：")
        for w in warnings:
            print("  -", w)
    if errors:
        print("\n错误：")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    print("\n全部通过：文档 / 任务书 / 课件 / 链接 / 同步状态正常。")


if __name__ == "__main__":
    main()
