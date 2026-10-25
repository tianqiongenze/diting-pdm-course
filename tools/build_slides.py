"""课程课件生成器（16 套 PPT，含讲稿备注）。

设计原则：
1. 统一视觉规范（深蓝工业风 + 橙色强调），不依赖模板文件；
2. 每页信息量受控（≤5 条要点，每条 ≤2 行），避免上课"念 PPT"；
3. 每页写讲稿备注（notes），教师可直接照读；
4. 内容与 03_课程教学方案 / 09_项目任务书 一一对应。

用法：
    python tools/build_slides.py            # 全部重建
    python tools/build_slides.py --only 7   # 只重建第 7 讲
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "slides" / "pptx"
sys.path.insert(0, str(Path(__file__).resolve().parent / "src"))

FONT = "微软雅黑"
MONO = "Consolas"

INK = RGBColor(0x0E, 0x1B, 0x2A)
INK2 = RGBColor(0x18, 0x2B, 0x40)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xF2, 0xF5, 0xF7)
LINE = RGBColor(0xD8, 0xDF, 0xE6)
TEXT = RGBColor(0x1B, 0x27, 0x33)
SUB = RGBColor(0x5A, 0x6B, 0x7C)
ACCENT = RGBColor(0xE0, 0x4F, 0x2B)
ACCENT2 = RGBColor(0x1F, 0x7A, 0x8C)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DIM = RGBColor(0x8A, 0x9A, 0xAB)

W, H = 13.333, 7.5
MARGIN = 0.85


def _set_font(run, name: str = FONT) -> None:
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def textbox(slide, x, y, w, h, text="", size=18, color=TEXT, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
            line_spacing=1.25, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    _set_font(r, font)
    return tb


def rect(slide, x, y, w, h, fill, line=None, radius=None, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = shadow
    style = s._element.find(qn("p:style"))
    if style is not None:
        s._element.remove(style)
    return s


def bullets(slide, x, y, w, h, items, size=19, color=TEXT, gap=10, bullet_color=ACCENT,
            line_spacing=1.3):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap)
        if isinstance(item, tuple):
            head, detail = item
            r0 = p.add_run()
            r0.text = "▪  "
            r0.font.size = Pt(size - 2)
            r0.font.color.rgb = bullet_color
            _set_font(r0)
            r1 = p.add_run()
            r1.text = head
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = color
            _set_font(r1)
            if detail:
                r2 = p.add_run()
                r2.text = "  " + detail
                r2.font.size = Pt(size - 2)
                r2.font.color.rgb = SUB
                _set_font(r2)
        else:
            r0 = p.add_run()
            r0.text = "▪  "
            r0.font.size = Pt(size - 2)
            r0.font.color.rgb = bullet_color
            _set_font(r0)
            r1 = p.add_run()
            r1.text = item
            r1.font.size = Pt(size)
            r1.font.color.rgb = color
            _set_font(r1)
    return tb


def header(slide, kicker, title, dark=False):
    bg = INK if dark else PAPER
    rect(slide, 0, 0, W, H, bg)
    fg = WHITE if dark else TEXT
    textbox(slide, MARGIN, 0.52, 8.5, 0.3, kicker, 12, ACCENT if not dark else GOLD, True)
    textbox(slide, MARGIN, 0.86, 11.4, 0.75, title, 30, fg, True)
    rect(slide, MARGIN, 1.72, 1.1, 0.055, ACCENT)
    if not dark:
        rect(slide, 0, 0, 0.14, H, ACCENT)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- layouts

def slide_cover(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, INK)
    # 大号讲次数字
    textbox(s, 10.15, 0.35, 3.0, 2.4, f"{d['no']:02d}", 120, RGBColor(0x1E, 0x33, 0x4A), True)
    rect(s, 0, 0, 0.18, H, ACCENT)
    textbox(s, MARGIN, 1.45, 8.0, 0.35, f"{d['module']}  ·  第 {d['no']:02d} 讲", 14, GOLD, True)
    textbox(s, MARGIN, 2.0, 9.1, 1.9, d["title"], 42, WHITE, True, line_spacing=1.15)
    textbox(s, MARGIN, 4.15, 8.3, 0.9, d["subtitle"], 17, RGBColor(0xB8, 0xC6, 0xD4), line_spacing=1.4)
    rect(s, MARGIN, 5.35, 8.2, 0.02, RGBColor(0x33, 0x48, 0x5E))
    textbox(s, MARGIN, 5.6, 9.2, 0.35, d["meta"], 13, DIM)
    textbox(s, MARGIN, 6.55, 8.0, 0.35, "AI 智能体工程化项目实战  ·  64 学时  ·  双元制本土化课程", 11,
            RGBColor(0x6F, 0x82, 0x96))
    notes(s, d.get("cover_notes", f"本讲为第 {d['no']:02d} 讲，{d['title']}。开场先讲情境，不要先讲概念。"))
    return s


def slide_nav(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "SESSION MAP", "本讲导航：4 学时怎么用")
    cols = d["agenda"]
    w = 2.85
    for i, (t, sub) in enumerate(cols):
        x = MARGIN + i * (w + 0.22)
        rect(s, x, 2.15, w, 1.75, MUTED, radius=0.06)
        rect(s, x, 2.15, 0.06, 1.75, ACCENT if i % 2 == 0 else ACCENT2)
        textbox(s, x + 0.28, 2.38, w - 0.5, 0.35, f"0{i + 1}", 13, ACCENT if i % 2 == 0 else ACCENT2, True)
        textbox(s, x + 0.28, 2.72, w - 0.5, 0.4, t, 17, TEXT, True)
        textbox(s, x + 0.28, 3.15, w - 0.5, 0.6, sub, 12, SUB, line_spacing=1.25)
    # 六步
    steps = ["信息", "计划", "决策", "实施", "检查", "评价"]
    y = 4.85
    textbox(s, MARGIN, 4.35, 4.0, 0.3, "行动六步（完整行动模型）", 12, SUB, True)
    for i, st in enumerate(steps):
        x = MARGIN + i * 1.95
        rect(s, x, y, 1.62, 0.62, INK if i % 2 == 0 else INK2, radius=0.12)
        textbox(s, x, y + 0.17, 1.62, 0.3, st, 14, WHITE, True, PP_ALIGN.CENTER)
        if i < 5:
            textbox(s, x + 1.64, y + 0.16, 0.3, 0.3, "›", 16, DIM, True, PP_ALIGN.CENTER)
    textbox(s, MARGIN, 5.85, 11.5, 0.8,
            "每一讲都走完这六步：接到任务 → 想清楚 → 定方案 → 做出来 → 自己验 → 被评价。",
            14, SUB)
    notes(s, "先让学生看到节奏：4 学时不是 4 节理论课，而是 1 个完整工作循环。强调最后 25 分钟验收不可省。")
    return s


def slide_objectives(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "LEARNING OBJECTIVES", "本讲目标：3 知识 + 3 技能 + 1 素养")
    groups = [("知识", d["obj"]["k"], ACCENT2), ("技能", d["obj"]["s"], ACCENT), ("素养", d["obj"]["a"], GOLD)]
    w = 3.62
    for i, (name, items, color) in enumerate(groups):
        x = MARGIN + i * (w + 0.35)
        rect(s, x, 2.1, w, 0.5, color)
        textbox(s, x + 0.22, 2.2, w - 0.4, 0.3, name, 15, WHITE, True)
        rect(s, x, 2.6, w, 3.1, MUTED)
        for j, it in enumerate(items):
            textbox(s, x + 0.25, 2.82 + j * 0.95, w - 0.5, 0.85, f"{j + 1}. {it}", 13, TEXT, line_spacing=1.3)
    textbox(s, MARGIN, 6.05, 11.6, 0.5, "达标标准：能做出工件 + 能说清依据 + 能通过交叉验收。", 14, SUB)
    notes(s, "逐条读目标，但不要解释。告诉学生：目标不是背的，是下课前用来对照的。")
    return s


def slide_context(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, INK)
    rect(s, 0, 0, 0.18, H, ACCENT)
    textbox(s, MARGIN, 0.75, 6.0, 0.3, "SITUATION · 现场", 12, GOLD, True)
    textbox(s, MARGIN, 1.55, 11.2, 3.5, d["context"], 21, WHITE, line_spacing=1.6, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, MARGIN, 5.35, 5.4, 0.02, RGBColor(0x33, 0x48, 0x5E))
    textbox(s, MARGIN, 5.55, 5.4, 0.35, f"委托方：{d['client']}", 13, RGBColor(0xB8, 0xC6, 0xD4))
    textbox(s, MARGIN, 5.95, 11.2, 0.7, f"核心矛盾：{d['conflict']}", 15, GOLD, True)
    notes(s, "讲情境时不要急着给答案。讲完停 5 秒，问：'如果是你，你怎么回答？' 让 2–3 个学生说。")
    return s


def slide_questions(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "QUESTION CHAIN", "问题链：今天要回答的 5 个问题")
    qs = d["questions"]
    for i, q in enumerate(qs):
        y = 2.0 + i * 0.92
        textbox(s, MARGIN, y, 0.7, 0.5, f"{i + 1:02d}", 20, ACCENT, True)
        textbox(s, MARGIN + 0.75, y - 0.02, 10.6, 0.7, q, 17, TEXT, line_spacing=1.25)
        if i < len(qs) - 1:
            rect(s, MARGIN + 0.75, y + 0.7, 10.4, 0.008, LINE)
    textbox(s, MARGIN, 6.75, 11.0, 0.35, "下课前，这 5 个问题你要能自己回答，而不是背答案。", 13, SUB)
    notes(s, "问题链是本讲的骨架。每讲完一个概念，回到对应问题，确认学生能答。")
    return s


def slide_bullets(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, spec.get("kicker", d["module"]), spec["title"])
    n = len(spec["items"])
    gap = 30 if n <= 4 else (24 if n == 5 else 15)
    bullets(s, MARGIN, 2.0, 11.4, 4.4, spec["items"], size=spec.get("size", 21), gap=gap)
    if spec.get("foot"):
        textbox(s, MARGIN, 6.55, 11.4, 0.5, spec["foot"], 13, SUB)
    notes(s, spec.get("notes", ""))
    return s


def slide_two_col(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, spec.get("kicker", d["module"]), spec["title"])
    left, right = spec["left"], spec["right"]
    for i, (col, x, color) in enumerate([(left, MARGIN, ACCENT2), (right, 7.0, ACCENT)]):
        rect(s, x, 2.05, 5.45, 0.06, color)
        textbox(s, x, 2.2, 5.45, 0.4, col["head"], 16, color, True)
        bullets(s, x, 2.8, 5.35, 3.5, col["items"], size=spec.get("size", 18), gap=24,
                bullet_color=color)
    if spec.get("foot"):
        textbox(s, MARGIN, 6.5, 11.4, 0.4, spec["foot"], 13, SUB)
    notes(s, spec.get("notes", ""))
    return s


def slide_table(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, spec.get("kicker", d["module"]), spec["title"])
    rows = len(spec["data"])
    cols = len(spec["data"][0])
    tbl_w = 11.6
    shape = s.shapes.add_table(rows, cols, Inches(MARGIN), Inches(2.05), Inches(tbl_w), Inches(0.5 * rows))
    table = shape.table
    widths = spec.get("widths") or [1] * cols
    total = sum(widths)
    for i, wd in enumerate(widths):
        table.columns[i].width = Emu(int(Inches(tbl_w) * wd / total))
    for r, row in enumerate(spec["data"]):
        table.rows[r].height = Inches(spec.get("row_h", 0.52))
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = INK
            else:
                cell.fill.fore_color.rgb = PAPER if r % 2 else MUTED
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.1
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(spec.get("size", 14))
            run.font.bold = r == 0
            run.font.color.rgb = WHITE if r == 0 else TEXT
            _set_font(run)
    if spec.get("foot"):
        textbox(s, MARGIN, 6.7, 11.4, 0.4, spec["foot"], 12, SUB)
    notes(s, spec.get("notes", ""))
    return s


def slide_code(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, spec.get("kicker", d["module"]), spec["title"])
    rect(s, MARGIN, 2.05, 11.6, 4.25, INK, radius=0.02)
    textbox(s, MARGIN + 0.3, 2.18, 6.0, 0.3, spec.get("file", ""), 11, RGBColor(0x8F, 0xA3, 0xB6), False, font=MONO)
    lines = spec["code"].split("\n")
    tb = s.shapes.add_textbox(Inches(MARGIN + 0.3), Inches(2.55), Inches(11.0), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = False
    for i, line in enumerate(lines[:17]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.18
        r = p.add_run()
        r.text = line if line else " "
        r.font.size = Pt(spec.get("size", 13))
        r.font.color.rgb = RGBColor(0xD6, 0xE2, 0xEE)
        _set_font(r, MONO)
    if spec.get("foot"):
        textbox(s, MARGIN, 6.55, 11.4, 0.5, spec["foot"], 12, SUB)
    notes(s, spec.get("notes", ""))
    return s


def slide_steps(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, spec.get("kicker", d["module"]), spec["title"])
    items = spec["items"]
    n = len(items)
    w = (11.6 - 0.2 * (n - 1)) / n
    for i, (head, detail) in enumerate(items):
        x = MARGIN + i * (w + 0.2)
        color = ACCENT if i % 2 == 0 else ACCENT2
        rect(s, x, 2.15, w, 0.09, color)
        textbox(s, x, 2.35, w, 0.3, f"{i + 1}", 13, color, True)
        textbox(s, x, 2.65, w, 0.55, head, 14, TEXT, True, line_spacing=1.2)
        textbox(s, x, 3.18, w, 1.2, detail, 13, SUB, line_spacing=1.4)
    if spec.get("foot"):
        rect(s, MARGIN, 4.65, 11.6, 1.15, MUTED, radius=0.04)
        textbox(s, MARGIN + 0.3, 4.9, 11.0, 0.75, spec["foot"], 14, TEXT, line_spacing=1.45)
    notes(s, spec.get("notes", ""))
    return s


def slide_checklist(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "DEFINITION OF DONE", spec.get("title", "检查表：做到这些才算完成"))
    items = spec["items"]
    half = (len(items) + 1) // 2
    for col in range(2):
        chunk = items[col * half:(col + 1) * half]
        for i, it in enumerate(chunk):
            y = 2.1 + i * (3.9 / max(half, 1))
            x = MARGIN + col * 5.9
            rect(s, x, y + 0.06, 0.22, 0.22, PAPER, LINE)
            textbox(s, x + 0.4, y, 5.2, 0.65, it, 15, TEXT, line_spacing=1.3)
    textbox(s, MARGIN, 6.55, 11.4, 0.4, "自检 → 换组交叉验收 → 双方签字，缺一项不算完成。", 13, SUB)
    notes(s, spec.get("notes", "让学生逐项打勾。交叉验收时，验收方要能指出证据位置。"))
    return s


def slide_rubric(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "ASSESSMENT", spec.get("title", "本讲评分点"))
    data = [["评分点", "权重", "评分人"]] + spec["data"]
    rows, cols = len(data), 3
    shape = s.shapes.add_table(rows, cols, Inches(MARGIN), Inches(2.15), Inches(9.4), Inches(0.5 * rows))
    table = shape.table
    table.columns[0].width = Inches(5.0)
    table.columns[1].width = Inches(1.9)
    table.columns[2].width = Inches(2.5)
    for r, row in enumerate(data):
        table.rows[r].height = Inches(0.52)
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.14)
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK if r == 0 else (PAPER if r % 2 else MUTED)
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(13)
            run.font.bold = r == 0
            run.font.color.rgb = WHITE if r == 0 else TEXT
            _set_font(run)
    textbox(s, 10.45, 2.2, 2.4, 1.2, "总分", 13, SUB, True)
    textbox(s, 10.4, 2.5, 2.6, 1.0, f"{spec['total']}分", 38, ACCENT, True)
    textbox(s, 10.45, 3.6, 2.4, 1.6, "计入\n工件作品集", 12, SUB, line_spacing=1.5)
    if spec.get("foot"):
        textbox(s, MARGIN, 6.55, 11.4, 0.4, spec["foot"], 12, SUB)
    else:
        textbox(s, MARGIN, 5.6, 9.2, 1.2,
                "评分规则：按证据打分，不看口才；关键项缺失整项降级；红线违规一票否决。",
                12, SUB, line_spacing=1.5)
    notes(s, spec.get("notes", "打分看证据，不看口才。"))
    return s


def slide_summary(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "TAKEAWAYS", spec.get("title", "本讲小结：带走 4 句话"))
    items = spec["items"]
    for i, (head, detail) in enumerate(items):
        y = 2.05 + i * 1.12
        rect(s, MARGIN, y, 11.6, 0.98, MUTED, radius=0.05)
        rect(s, MARGIN, y, 0.07, 0.98, ACCENT if i % 2 == 0 else ACCENT2)
        textbox(s, MARGIN + 0.32, y + 0.13, 3.6, 0.35, head, 16, TEXT, True)
        textbox(s, MARGIN + 0.32, y + 0.52, 10.9, 0.4, detail, 13, SUB)
    notes(s, spec.get("notes", "小结不要重复讲过的内容，只留 4 句能带走的话。"))
    return s


def slide_homework(prs, d, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "AFTER CLASS", "课后任务与预习")
    bullets(s, MARGIN, 2.1, 7.2, 3.8, spec["items"], size=17)
    rect(s, 8.55, 2.1, 3.95, 3.9, MUTED, radius=0.05)
    textbox(s, 8.85, 2.35, 3.4, 0.35, "提交要求", 14, ACCENT, True)
    textbox(s, 8.85, 2.8, 3.4, 2.9, spec.get("submit", "位置：artifacts/T{NN}/\n截止：课后 48 小时\n必写：README（做了什么/怎么验证/遗留）"),
            12, TEXT, line_spacing=1.5)
    notes(s, spec.get("notes", "作业要当堂说清楚提交位置与截止时间。"))
    return s


def slide_end(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, INK)
    rect(s, 0, 0, 0.18, H, ACCENT)
    textbox(s, MARGIN, 1.6, 8.0, 0.35, "NEXT", 13, GOLD, True)
    textbox(s, MARGIN, 2.0, 9.6, 0.9, d["next"], 34, WHITE, True, line_spacing=1.2)
    textbox(s, MARGIN, 3.1, 9.6, 0.8, d.get("next_hint", ""), 15, RGBColor(0xB8, 0xC6, 0xD4), line_spacing=1.4)
    rect(s, MARGIN, 4.55, 4.6, 0.02, RGBColor(0x33, 0x48, 0x5E))
    textbox(s, MARGIN, 4.8, 10.5, 1.2, "下课前 10 分钟：复盘三问\n① 凭什么说它是对的？ ② 明天最先坏在哪？ ③ 哪一步可以少走？",
            14, RGBColor(0x9F, 0xB2, 0xC4), line_spacing=1.7)
    notes(s, "结束页不要拖堂。留 10 分钟做复盘三问，比多讲一个知识点有价值。")
    return s


RENDERERS = {
    "bullets": slide_bullets,
    "two_col": slide_two_col,
    "table": slide_table,
    "code": slide_code,
    "steps": slide_steps,
    "checklist": slide_checklist,
    "rubric": slide_rubric,
    "summary": slide_summary,
    "homework": slide_homework,
}


def build_deck(d) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide_cover(prs, d)
    slide_nav(prs, d)
    slide_objectives(prs, d)
    slide_context(prs, d)
    slide_questions(prs, d)
    for spec in d["slides"]:
        fn = RENDERERS[spec["type"]]
        if spec["type"] in ("bullets", "two_col", "table", "code", "steps", "checklist", "rubric", "summary", "homework"):
            fn(prs, d, spec)
    slide_end(prs, d)

    OUT.mkdir(parents=True, exist_ok=True)
    safe = d["title"].replace("/", "·").replace("：", "-").replace(" ", "")
    path = OUT / f"第{d['no']:02d}讲_{safe}.pptx"
    prs.save(path)
    return path


def load_decks():
    decks = []
    for mod in ("content_m1", "content_m2", "content_m3", "content_m4"):
        m = __import__(mod)
        decks.extend(m.DECKS)
    return sorted(decks, key=lambda x: x["no"])


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", type=int, default=0)
    ap.add_argument("--changed", action="store_true", help="仅重建源文件更新过的讲次（简化为全部重建）")
    args = ap.parse_args()
    decks = load_decks()
    assert len(decks) == 16, f"expected 16 decks, got {len(decks)}"
    for d in decks:
        if args.only and d["no"] != args.only:
            continue
        p = build_deck(d)
        print(f"built: {p.name}  ({len(p.name)} chars)")
    print("done")


if __name__ == "__main__":
    main()
